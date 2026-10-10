"""Export helpers: Excel (.xlsx) and PDF for complaint reports.

Reuses the Arabic font (Noto Naskh Arabic) and the reshape/bidi pipeline from
the existing maintenance system so Arabic renders correctly.
"""

import io
from datetime import datetime
from pathlib import Path

import arabic_reshaper
from bidi.algorithm import get_display
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Image,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from . import database as db

BASE_DIR = Path(__file__).resolve().parent.parent
FONT_PATH = BASE_DIR / "static" / "fonts" / "NotoNaskhArabic.ttf"
LOGO_PATH = BASE_DIR / "static" / "img" / "jazan-cluster-logo.png"

HOSPITAL_AR = "مستشفى الأمير محمد بن ناصر — جازان"
HOSPITAL_EN = "PRINCE MOHAMMED BIN NASER HOSPITAL — MAINTENANCE & OPERATIONS"

STATUS_AR = {s: ar for s, ar, _ in db.STATUSES}
STATUS_AR[db.CANCELLED] = "ملغي"
PRIORITY_AR = {p: ar for p, ar, _ in db.PRIORITIES}


def _ar(text) -> str:
    t = str(text or "")
    if not t:
        return ""
    try:
        return get_display(arabic_reshaper.reshape(t))
    except Exception:
        return t


def _wrap(text) -> str:
    t = str(text or "")
    return (t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))


def _cell(text, bold=False, size=8):
    val = _wrap(_ar(text))
    if bold:
        val = f"<b>{val}</b>"
    return Paragraph(val, ParagraphStyle(
        "cell", fontName="NaskhBold" if bold else "Naskh", fontSize=size, leading=size + 3,
    ))


def _label(text):
    return Paragraph(_wrap(_ar(text)), ParagraphStyle("l", fontName="NaskhBold", fontSize=9, leading=12))


def _fmt(v):
    if not v:
        return "-"
    try:
        d = datetime.fromisoformat(str(v))
        return f"{d.year}/{d.month:02d}/{d.day:02d} {d.hour:02d}:{d.minute:02d}"
    except Exception:
        return str(v)


def _row_meta(c):
    return {
        "ticket_no": c["ticket_no"],
        "title": c["title"] or c["description"][:60],
        "department": c["department"],
        "location": c["location_text"] or "-",
        "fault_type": c["fault_type"],
        "specialty": c["specialty"] or "-",
        "priority": PRIORITY_AR.get(c["priority"], c["priority"]),
        "safety": "نعم" if c["affects_safety"] else "لا",
        "status": STATUS_AR.get(c["status"], c["status"]),
        "technician": c["assigned_technician"] or "-",
        "created": _fmt(c["created_at"]),
        "description": c["description"],
        "reporter": c["reporter_name"],
        "contact": c["reporter_contact"],
    }


def _export_rows(filters=None):
    return [_row_meta(c) for c in db.report_rows(filters)]


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

def build_excel(filters=None, title="تقرير بلاغات الصيانة والتشغيل") -> bytes:
    rows = _export_rows(filters)
    wb = Workbook()
    ws = wb.active
    ws.title = "بلاغات الصيانة"

    fill = PatternFill("solid", fgColor="0B6B42")
    white = Font(color="FFFFFF", bold=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="CBD5E1")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.merge_cells("A1:L1")
    ws["A1"] = f"{HOSPITAL_AR} — {title}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = center
    ws.row_dimensions[1].height = 30

    ws.merge_cells("A2:L2")
    ws["A2"] = "قسم الصيانة والتشغيل | بتاريخ " + datetime.now().strftime("%Y/%m/%d %H:%M")
    ws["A2"].alignment = center
    ws["A2"].font = Font(size=9, color="64748B")

    headers = ["م", "رقم البلاغ", "العنوان", "القسم", "الموقع", "نوع العطل",
               "التخصص", "الأولوية", "سلامة المرضى", "الحالة", "الفني", "تاريخ الإرسال"]
    hr = 4
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=hr, column=i, value=h)
        c.fill = fill
        c.font = white
        c.alignment = center
        c.border = border

    r = hr + 1
    for idx, m in enumerate(rows, start=1):
        vals = [idx, m["ticket_no"], m["title"], m["department"], m["location"],
                m["fault_type"], m["specialty"], m["priority"], m["safety"],
                m["status"], m["technician"], m["created"]]
        for i, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=i, value=v)
            c.border = border
            c.alignment = Alignment(vertical="top", wrap_text=(i in (3, 4, 5)))
        r += 1

    widths = [5, 15, 28, 20, 22, 18, 16, 10, 12, 14, 14, 18]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

FONT = "Naskh"
FONT_B = "NaskhBold"


def _reg_font():
    global FONT, FONT_B
    try:
        pdfmetrics.registerFont(TTFont("Naskh", str(FONT_PATH)))
        bold = FONT_PATH.with_name("NotoNaskhArabic-Bold.ttf")
        if bold.exists():
            pdfmetrics.registerFont(TTFont("NaskhBold", str(bold)))
        else:
            pdfmetrics.registerFont(TTFont("NaskhBold", str(FONT_PATH)))
        pdfmetrics.registerFontFamily("Naskh", normal="Naskh", bold="NaskhBold",
                                      italic="Naskh", boldItalic="NaskhBold")
    except Exception:
        FONT, FONT_B = "Helvetica", "Helvetica-Bold"


_reg_font()


def build_pdf(filters=None, title="تقرير بلاغات الصيانة والتشغيل") -> bytes:
    rows = _export_rows(filters)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=landscape(A4),
        leftMargin=12 * mm, rightMargin=12 * mm,
        topMargin=12 * mm, bottomMargin=14 * mm,
        title=title,
    )

    p_title = ParagraphStyle("pT", fontName=FONT_B, fontSize=15, leading=20, alignment=TA_CENTER)
    p_sub = ParagraphStyle("pS", fontName=FONT, fontSize=8, leading=11, alignment=TA_CENTER,
                           textColor=colors.HexColor("#64748b"))

    story = []
    lh_data = [[
        "",
        Paragraph(
            f"<b>{_wrap(_ar(HOSPITAL_AR))}</b><br/>"
            f"<font name='{FONT}' size='7' color='#334155'>{_wrap(HOSPITAL_EN)}</font>",
            ParagraphStyle("lhc", fontName=FONT_B, fontSize=11, leading=15, alignment=TA_CENTER),
        ),
        "",
    ]]
    lh = Table(lh_data, colWidths=[30 * mm, 140 * mm, 30 * mm])
    if LOGO_PATH.exists():
        lh = Table(
            [[Image(str(LOGO_PATH), width=20 * mm, height=20 * mm),
              lh_data[0][1], Image(str(LOGO_PATH), width=20 * mm, height=20 * mm)]],
            colWidths=[30 * mm, 140 * mm, 30 * mm],
        )
    lh.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
    story.append(lh)
    story.append(Spacer(1, 4 * mm))
    story.append(Paragraph(_wrap(_ar(title)), p_title))
    story.append(Paragraph(_wrap(_ar("MAINTENANCE & OPERATIONS COMPLAINT REPORT")), p_sub))
    story.append(Spacer(1, 4 * mm))

    header = [_label("#"), _label("رقم البلاغ"), _label("العنوان"), _label("القسم"),
              _label("نوع العطل"), _label("الأولوية"), _label("الحالة"), _label("الموقع"),
              _label("تاريخ الإرسال")]
    data = [header]
    for idx, m in enumerate(rows, start=1):
        data.append([
            _ar(str(idx)),
            _cell(m["ticket_no"], bold=True),
            _cell(m["title"]),
            _cell(m["department"]),
            _cell(m["fault_type"]),
            _ar(m["priority"]),
            _ar(m["status"]),
            _cell(m["location"]),
            _ar(m["created"]),
        ])
    table = Table(data, colWidths=[10 * mm, 28 * mm, 52 * mm, 30 * mm, 26 * mm,
                                   20 * mm, 24 * mm, 40 * mm, 30 * mm], repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0b6b42")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#cbd5e1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("ALIGN", (0, 0), (0, -1), "CENTER"),
    ]))
    story.append(table)
    story.append(Spacer(1, 4 * mm))
    story.append(Paragraph(
        _wrap(_ar(f"{HOSPITAL_AR} — {title} | Generated: {datetime.now().strftime('%Y/%m/%d')}")),
        ParagraphStyle("ft", fontName=FONT, fontSize=7, leading=10, alignment=TA_CENTER,
                       textColor=colors.HexColor("#64748b")),
    ))
    doc.build(story)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# QR codes
# ---------------------------------------------------------------------------

def make_qr_png(data: str, box_size: int = 7, border: int = 2) -> bytes:
    """Return a PNG QR code (bytes) for the given text/URL."""
    import qrcode
    from qrcode.constants import ERROR_CORRECT_M

    qr = qrcode.QRCode(version=None, error_correction=ERROR_CORRECT_M,
                       box_size=box_size, border=border)
    qr.add_data(data)
    qr.make(fit=True)
    img = qr.make_image(fill_color="black", back_color="white")
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Single-complaint PDF (printable card with QR + full tracking link)
# ---------------------------------------------------------------------------

def build_complaint_pdf(c, track_url: str, events=None) -> bytes:
    """A one-page printable complaint card (Arabic) including a QR code that
    links to the public tracking page."""
    m = _row_meta(c)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4,
        leftMargin=15 * mm, rightMargin=15 * mm,
        topMargin=14 * mm, bottomMargin=14 * mm,
        title=f"Complaint {c['ticket_no']}",
    )

    p_title = ParagraphStyle("ct", fontName=FONT_B, fontSize=15, leading=20, alignment=TA_CENTER)
    p_sub = ParagraphStyle("cs", fontName=FONT, fontSize=8, leading=11, alignment=TA_CENTER,
                           textColor=colors.HexColor("#64748b"))

    story = []
    lh_data = [[
        Image(str(LOGO_PATH), width=18 * mm, height=18 * mm) if LOGO_PATH.exists() else "",
        Paragraph(
            f"<b>{_wrap(_ar(HOSPITAL_AR))}</b><br/>"
            f"<font name='{FONT}' size='7' color='#334155'>{_wrap(HOSPITAL_EN)}</font>",
            ParagraphStyle("lh", fontName=FONT_B, fontSize=11, leading=15, alignment=TA_CENTER),
        ),
        Image(str(LOGO_PATH), width=18 * mm, height=18 * mm) if LOGO_PATH.exists() else "",
    ]]
    lh = Table(lh_data, colWidths=[25 * mm, 130 * mm, 25 * mm])
    lh.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
    story.append(lh)
    story.append(Spacer(1, 5 * mm))
    story.append(Paragraph(_wrap(_ar("بطاقة بلاغ صيانة")), p_title))
    story.append(Paragraph(_wrap(_ar("MAINTENANCE COMPLAINT CARD")), p_sub))
    story.append(Spacer(1, 5 * mm))

    qr_img = Image(io.BytesIO(make_qr_png(track_url, box_size=8, border=2)),
                   width=34 * mm, height=34 * mm)
    qr_cell = [
        qr_img,
        Paragraph(_wrap(_ar("امسح الرمز للمتابعة")),
                  ParagraphStyle("qc", fontName=FONT, fontSize=7, leading=9, alignment=TA_CENTER)),
        Paragraph(_wrap(track_url),
                  ParagraphStyle("qu", fontName=FONT, fontSize=6.5, leading=8, alignment=TA_CENTER,
                                 textColor=colors.HexColor("#0b6b42"))),
    ]

    def kv(label, value, bold=False):
        return [_label(label), _cell(value, bold=bold)]

    rows = [
        kv("رقم البلاغ", c["ticket_no"], bold=True),
        kv("العنوان", m["title"]),
        kv("القسم", m["department"]),
        kv("الموقع", m["location"]),
        kv("نوع العطل", m["fault_type"]),
        kv("التخصص", m["specialty"]),
        kv("الأولوية", m["priority"]),
        kv("الحالة", m["status"]),
        kv("سلامة المرضى", m["safety"]),
        kv("الفني المسؤول", m["technician"]),
        kv("مقدّم البلاغ", f"{m['reporter']} — {m['contact']}"),
        kv("تاريخ الإرسال", m["created"]),
        kv("الوصف", m["description"]),
    ]
    # Pair the QR cell beside the first few rows
    info = Table(rows, colWidths=[28 * mm, 85 * mm])
    info.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#cbd5e1")),
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#ecfdf5")),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    qr_table = Table([[qr_cell]], colWidths=[42 * mm])
    qr_table.setStyle(TableStyle([
        ("BOX", (0, 0), (-1, -1), 0.4, colors.HexColor("#0b6b42")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    outer = Table([[info, qr_table]], colWidths=[113 * mm, 42 * mm])
    outer.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story.append(outer)

    if events:
        story.append(Spacer(1, 5 * mm))
        story.append(_label("سجل المتابعة"))
        ev_rows = [[_label("الإجراء"), _label("التفاصيل"), _label("التاريخ")]]
        for e in events[-15:]:
            ev_rows.append([
                _cell(e["action"]),
                _cell((str(e["to_value"]) if e["to_value"] else (e["notes"] or ""))),
                _ar(_fmt(e["created_at"])),
            ])
        ev = Table(ev_rows, colWidths=[50 * mm, 75 * mm, 30 * mm], repeatRows=1)
        ev.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0b6b42")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#cbd5e1")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]))
        story.append(ev)

    story.append(Spacer(1, 4 * mm))
    story.append(Paragraph(
        _wrap(_ar(f"{HOSPITAL_AR} | تم التحديث بتاريخ {datetime.now().strftime('%Y/%m/%d %H:%M')}")),
        ParagraphStyle("cf", fontName=FONT, fontSize=7, leading=10, alignment=TA_CENTER,
                       textColor=colors.HexColor("#64748b")),
    ))
    doc.build(story)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Live hierarchical Excel (grouped by department) - auto-saved on every report
# ---------------------------------------------------------------------------

def build_live_excel(filters=None) -> bytes:
    """Workbook with two sheets:
    * "شجري"  - tree view grouped by department (collapsible outline groups).
    * "كامل"  - flat table with every column.
    Always reflects the current database contents.
    """
    rows = [_row_meta(c) for c in db.report_rows(filters or {})]
    wb = Workbook()

    fill = PatternFill("solid", fgColor="0B6B42")
    dept_fill = PatternFill("solid", fgColor="DCFCE7")
    white = Font(color="FFFFFF", bold=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="CBD5E1")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # --- Tree sheet -------------------------------------------------------
    ws = wb.active
    ws.title = "شجري"
    ws.merge_cells("A1:H1")
    ws["A1"] = f"{HOSPITAL_AR} — شجرة بلاغات الصيانة"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = center
    ws.row_dimensions[1].height = 28
    ws.merge_cells("A2:H2")
    ws["A2"] = "قسم الصيانة والتشغيل | آخر تحديث " + datetime.now().strftime("%Y/%m/%d %H:%M")
    ws["A2"].alignment = center
    ws["A2"].font = Font(size=9, color="64748B")

    tree_headers = ["القسم / البلاغ", "رقم البلاغ", "نوع العطل", "الأولوية",
                    "الحالة", "الفني", "الموقع", "تاريخ الإرسال"]
    for i, h in enumerate(tree_headers, start=1):
        c = ws.cell(row=4, column=i, value=h)
        c.fill = fill; c.font = white; c.alignment = center; c.border = border

    by_dept = {}
    for m in rows:
        by_dept.setdefault(m["department"], []).append(m)

    r = 5
    for dept in sorted(by_dept):
        items = by_dept[dept]
        dc = ws.cell(row=r, column=1, value=f"🏢 {dept}  ({len(items)})")
        dc.font = Font(bold=True, color="0B6B42")
        dc.fill = dept_fill
        for i in range(1, 9):
            ws.cell(row=r, column=i).border = border
            ws.cell(row=r, column=i).fill = dept_fill
        r += 1
        start = r
        for m in items:
            vals = [m["title"] or m["ticket_no"], m["ticket_no"], m["fault_type"],
                    m["priority"], m["status"], m["technician"], m["location"], m["created"]]
            for i, v in enumerate(vals, start=1):
                cc = ws.cell(row=r, column=i, value=v)
                cc.border = border
                cc.alignment = Alignment(vertical="top", wrap_text=(i in (1, 3, 7)))
            r += 1
        # Collapsible outline for the department's detail rows
        if r - 1 >= start:
            ws.row_dimensions.group(start, r - 1, outline_level=1)
    for i, w in enumerate([30, 15, 20, 10, 14, 14, 22, 18], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # --- Flat sheet -------------------------------------------------------
    ws2 = wb.create_sheet("كامل")
    headers = ["م", "رقم البلاغ", "العنوان", "القسم", "الموقع", "نوع العطل", "التخصص",
               "الأولوية", "سلامة المرضى", "الحالة", "الفني", "تاريخ الإرسال", "الوصف",
               "مقدّم البلاغ", "الجوال"]
    for i, h in enumerate(headers, start=1):
        c = ws2.cell(row=1, column=i, value=h)
        c.fill = fill; c.font = white; c.alignment = center; c.border = border
    rr = 2
    for idx, m in enumerate(rows, start=1):
        vals = [idx, m["ticket_no"], m["title"], m["department"], m["location"],
                m["fault_type"], m["specialty"], m["priority"], m["safety"], m["status"],
                m["technician"], m["created"], m["description"], m["reporter"], m["contact"]]
        for i, v in enumerate(vals, start=1):
            cc = ws2.cell(row=rr, column=i, value=v)
            cc.border = border
            cc.alignment = Alignment(vertical="top", wrap_text=(i in (3, 5, 13)))
        rr += 1
    for i, w in enumerate([5, 15, 26, 20, 22, 20, 16, 10, 12, 14, 14, 18, 40, 18, 16], start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()